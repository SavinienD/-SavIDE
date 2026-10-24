"""Extraction du texte des documents binaires pour ``files.read`` (outil interne).

Formats pris en charge : classeurs Excel (``.xlsx``/``.xlsm``, openpyxl),
documents Word (``.docx``, python-docx), PDF (pypdf), PowerPoint (``.pptx``),
texte/tableur/présentation OpenDocument (``.odt``/``.ods``/``.odp``) et epub
avec la bibliothèque standard. Les bibliothèques sont importées à la demande ;
quand l'une manque, l'erreur dit quel paquet installer avec ``shell``.
"""

from __future__ import annotations

import html
import os
import re
import zipfile
from typing import Callable, Optional

# Per-sheet / per-document guards so a huge file cannot flood the context.
MAX_SHEET_ROWS = 2000
MAX_CELL_CHARS = 500


def _missing(package: str, module: str) -> ImportError:
    return ImportError(
        f"lire ce format demande le paquet `{package}` "
        f"(nom d'import `{module}`) : installez-le avec "
        f"shell.run('python -m pip install {package}') puis réessayez"
    )


def _cell(value: object) -> str:
    if value is None:
        return ""
    text = str(value).replace("\r\n", " ").replace("\n", " ").replace("\t", " ")
    return text if len(text) <= MAX_CELL_CHARS else text[:MAX_CELL_CHARS] + "…"


def _xlsx(path: str) -> str:
    try:
        import openpyxl  # type: ignore
    except ImportError as err:
        raise _missing("openpyxl", "openpyxl") from err
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out: list[str] = []
    try:
        for sheet in book.worksheets:
            out.append(f"## Feuille : {sheet.title}")
            rows = 0
            for row in sheet.iter_rows(values_only=True):
                cells = [_cell(value) for value in row]
                while cells and not cells[-1]:
                    cells.pop()
                if not cells:
                    continue
                if rows >= MAX_SHEET_ROWS:
                    out.append(f"… (feuille tronquée après {MAX_SHEET_ROWS} lignes non vides)")
                    break
                out.append("\t".join(cells))
                rows += 1
            out.append("")
    finally:
        book.close()
    return "\n".join(out)


def _docx(path: str) -> str:
    try:
        import docx  # type: ignore
    except ImportError as err:
        raise _missing("python-docx", "docx") from err
    from docx.table import Table  # type: ignore
    from docx.text.paragraph import Paragraph  # type: ignore

    document = docx.Document(path)
    out: list[str] = []
    for block in document.element.body.iterchildren():
        tag = block.tag.rsplit("}", 1)[-1]
        if tag == "p":
            paragraph = Paragraph(block, document)
            text = paragraph.text.strip()
            if not text:
                continue
            style = (paragraph.style.name if paragraph.style is not None else "") or ""
            match = re.match(r"Heading (\d)", style) or re.match(r"Titre (\d)", style)
            if style == "Title":
                out.append(f"# {text}")
            elif match:
                out.append(f"{'#' * min(6, int(match.group(1)) + 1)} {text}")
            elif "List" in style or "Liste" in style:
                out.append(f"- {text}")
            else:
                out.append(text)
        elif tag == "tbl":
            for row in Table(block, document).rows:
                out.append(" | ".join(_cell(cell.text) for cell in row.cells))
            out.append("")
    return "\n".join(out)


def _pdf(path: str) -> str:
    try:
        import pypdf  # type: ignore
    except ImportError as err:
        raise _missing("pypdf", "pypdf") from err
    reader = pypdf.PdfReader(path)
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception as err:  # noqa: BLE001 - surface a clear message
            raise PermissionError(f"{os.path.basename(path)} est protégé par un mot de passe") from err
    out: list[str] = []
    for number, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        out.append(f"## Page {number}")
        out.append(text or "(aucun texte extractible sur cette page : image numérisée ?)")
        out.append("")
    return "\n".join(out)


_TAG = re.compile(r"<[^>]+>")
_BLOCK_END = re.compile(r"</(a:p|text:p|text:h|p|h[1-6]|li|tr|div|br)\s*>|<br\s*/?>|<text:line-break/>", re.I)


def _xml_text(xml: str) -> str:
    """Texte approximatif d'une partie XML/HTML : une ligne par bloc de type paragraphe."""
    xml = re.sub(r"<(script|style)\b.*?</\1>", "", xml, flags=re.S | re.I)
    xml = re.sub(r"<text:tab/>|<a:tab/>", "\t", xml)
    xml = _BLOCK_END.sub("\n", xml)
    text = html.unescape(_TAG.sub("", xml))
    lines = [re.sub(r"[ \t]+", " ", line).strip() for line in text.splitlines()]
    return "\n".join(line for line in lines if line)


def _slide_key(name: str) -> int:
    match = re.search(r"(\d+)\.xml$", name)
    return int(match.group(1)) if match else 0


def _pptx(path: str) -> str:
    out: list[str] = []
    with zipfile.ZipFile(path) as archive:
        slides = sorted(
            (n for n in archive.namelist() if re.match(r"ppt/slides/slide\d+\.xml$", n)),
            key=_slide_key,
        )
        for name in slides:
            out.append(f"## Diapositive {_slide_key(name)}")
            out.append(_xml_text(archive.read(name).decode("utf-8", "replace")))
            notes = f"ppt/notesSlides/notesSlide{_slide_key(name)}.xml"
            if notes in archive.namelist():
                text = _xml_text(archive.read(notes).decode("utf-8", "replace"))
                if text:
                    out.append(f"Notes : {text}")
            out.append("")
    return "\n".join(out)


def _opendocument(path: str) -> str:
    with zipfile.ZipFile(path) as archive:
        return _xml_text(archive.read("content.xml").decode("utf-8", "replace"))


def _epub(path: str) -> str:
    out: list[str] = []
    with zipfile.ZipFile(path) as archive:
        container = archive.read("META-INF/container.xml").decode("utf-8", "replace")
        opf_path = re.search(r'full-path="([^"]+)"', container)
        if not opf_path:
            raise ValueError("epub invalide : pas de paquet OPF")
        opf_name = opf_path.group(1)
        opf = archive.read(opf_name).decode("utf-8", "replace")
        base = opf_name.rsplit("/", 1)[0] + "/" if "/" in opf_name else ""
        manifest = {
            m.group(1): m.group(2)
            for m in re.finditer(r'<item\b[^>]*\bid="([^"]+)"[^>]*\bhref="([^"]+)"', opf)
        }
        manifest.update(
            {
                m.group(2): m.group(1)
                for m in re.finditer(r'<item\b[^>]*\bhref="([^"]+)"[^>]*\bid="([^"]+)"', opf)
            }
        )
        for idref in re.findall(r'<itemref\b[^>]*\bidref="([^"]+)"', opf):
            href = manifest.get(idref)
            if not href:
                continue
            name = os.path.normpath(base + href).replace(os.sep, "/")
            if name in archive.namelist():
                text = _xml_text(archive.read(name).decode("utf-8", "replace"))
                if text:
                    out.append(text)
                    out.append("")
    return "\n".join(out)


EXTRACTORS: dict[str, Callable[[str], str]] = {
    ".xlsx": _xlsx,
    ".xlsm": _xlsx,
    ".docx": _docx,
    ".docm": _docx,
    ".pdf": _pdf,
    ".pptx": _pptx,
    ".pptm": _pptx,
    ".odt": _opendocument,
    ".ods": _opendocument,
    ".odp": _opendocument,
    ".epub": _epub,
}

LEGACY = {
    ".doc": "Word 97-2003 (.doc)",
    ".xls": "Excel 97-2003 (.xls)",
    ".ppt": "PowerPoint 97-2003 (.ppt)",
}


def extractor_for(path: str) -> Optional[Callable[[str], str]]:
    return EXTRACTORS.get(os.path.splitext(path)[1].lower())


def legacy_format(path: str) -> Optional[str]:
    return LEGACY.get(os.path.splitext(path)[1].lower())


def looks_binary(data: bytes) -> bool:
    """Vrai pour un contenu qui n'est pas du texte (octets NUL dans les 8 premiers Kio)."""
    head = data[:8192]
    if head.startswith((b"\xff\xfe", b"\xfe\xff")):
        return False
    return b"\x00" in head
